import json
from pptx import Presentation

# Caminho do arquivo
pptx_path = r"template\template.pptx"

# Carregar apresentação
prs = Presentation(pptx_path)

# Função recursiva para navegar em formas agrupadas e aplicar filtro
def parse_shape(shape, slide_number, include_all=False):
    if not shape.shape_type:  # segurança
        return None

    # Ignorar conectores
    if "connector" in shape.name.lower():
        return None

    # Se for grupo
    if shape.shape_type == 6:  # msoGroup
        children = []
        for s in shape.shapes:
            child = parse_shape(s, slide_number, include_all)
            if child:
                children.append(child)
        if children or include_all:
            return {"name": shape.name, "children": children}
        else:
            return None

    # Se for gráfico
    elif shape.shape_type == 3 and hasattr(shape, "chart"):
        return {"name": shape.name, "type": "chart"}

    # Se for forma simples
    else:
        if include_all:
            return {"name": shape.name, "type": "shape"}
        else:
            nm = shape.name.lower()
            if nm == "filter" or nm.startswith("header") or nm.startswith("val") or nm.startswith("var"):
                return {"name": shape.name, "type": "shape"}
            else:
                return None

# Construir JSON
slides_json = {}

for i, slide in enumerate(prs.slides, start=1):
    slide_data = []
    include_all = (i == 1)  # Slide 1 pega tudo

    for shape in slide.shapes:
        element = parse_shape(shape, i, include_all)
        if element:
            slide_data.append(element)

    slides_json[f"Slide_{i}"] = {"page": i, "elements": slide_data}

# Salvar como JSON
output_path = "pptx_extracted.json"
with open(output_path, "w", encoding="utf-8") as f:
    json.dump(slides_json, f, ensure_ascii=False, indent=2)

print(f"Extração concluída! Arquivo salvo em: {output_path}")